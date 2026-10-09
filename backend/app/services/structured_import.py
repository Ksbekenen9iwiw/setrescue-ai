"""Preview-first import for simple tabular production XLSX files.

An uploaded spreadsheet is NOT automatically trusted; a Production Manager explicitly
applies its supported rows before they become structured project facts.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.errors import AppError
from app.models import Document, DocumentVersion, Project, Resource, Scene, Shot

MAX_IMPORT_ROWS = 500


def _cell(row: dict[str, Any], *names: str) -> str:
    for name in names:
        value = row.get(name)
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def _integer(value: str, default: int, minimum: int, maximum: int) -> int:
    if not value:
        return default
    try:
        parsed = int(float(value))
    except ValueError as exc:
        raise AppError(422, "DOCUMENT_PARSE_FAILED", f"Expected a number, got: {value[:60]}") from exc
    if not minimum <= parsed <= maximum:
        raise AppError(422, "DOCUMENT_PARSE_FAILED", f"Number must be between {minimum} and {maximum}.")
    return parsed


def preview(db: Session, project_id: str, document_id: str) -> dict[str, Any]:
    document = db.scalar(select(Document).where(Document.project_id == project_id, Document.id == document_id))
    if not document:
        raise AppError(404, "DOCUMENT_NOT_FOUND", "Document does not belong to the current project.")
    version = db.scalar(select(DocumentVersion).where(
        DocumentVersion.document_id == document.id,
        DocumentVersion.version_number == document.current_version,
    ))
    if not version or version.file_type.lower() != ".xlsx":
        raise AppError(422, "DOCUMENT_UNSUPPORTED", "Structured import currently supports XLSX; enter PDF facts manually.")
    source = Path(version.local_path)
    if not source.is_file():
        raise AppError(404, "DOCUMENT_NOT_FOUND", "Original local spreadsheet is unavailable.")
    wb = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    try:
        sheet = wb.worksheets[0]
        raw = sheet.iter_rows(values_only=True)
        try:
            header = next(raw)
        except StopIteration:
            raise AppError(422, "DOCUMENT_PARSE_FAILED", "The spreadsheet has no header row.")
        columns = [str(value or "").strip().lower().replace("_", " ") for value in header]
        rows = []
        for number, values in enumerate(raw, start=2):
            if number > MAX_IMPORT_ROWS + 1:
                raise AppError(422, "DOCUMENT_TOO_LARGE", "Structured import is limited to 500 rows per sheet.")
            if not any(value is not None and str(value).strip() for value in values):
                continue
            rows.append(dict(zip(columns, values)))
    finally:
        wb.close()

    items: list[dict[str, Any]] = []
    notes: list[str] = []
    kind = document.type
    for row in rows:
        if kind == "SCRIPT":
            number = _cell(row, "scene", "scene number", "scene no")
            title = _cell(row, "title", "scene title")
            if not number or not title:
                raise AppError(422, "DOCUMENT_PARSE_FAILED", "SCRIPT needs columns Scene and Title with nonempty values.")
            items.append({
                "scene_number": number[:20], "title": title[:160],
                "description": _cell(row, "description", "synopsis")[:5000],
                "estimated_duration_minutes": _integer(_cell(row, "minutes", "duration", "duration minutes"), 30, 1, 1440),
                "setup_minutes": _integer(_cell(row, "setup", "setup minutes"), 10, 0, 480),
            })
        elif kind == "SHOT_LIST":
            number = _cell(row, "scene", "scene number")
            shot_number = _cell(row, "shot", "shot number")
            if not number or not shot_number:
                raise AppError(422, "DOCUMENT_PARSE_FAILED", "SHOT_LIST needs Scene and Shot columns.")
            items.append({
                "scene_number": number[:20], "shot_number": shot_number[:30],
                "description": _cell(row, "description")[:2000],
                "frame_size": _cell(row, "frame", "frame size")[:40],
                "camera_movement": _cell(row, "movement", "camera movement")[:100],
                "estimated_duration_minutes": _integer(_cell(row, "minutes", "duration"), 10, 1, 1440),
            })
        elif kind == "RESOURCE_SHEET":
            name = _cell(row, "resource", "name", "resource name")
            type_ = _cell(row, "type", "resource type").upper()
            if not name or type_ not in {"CAST", "CREW", "EQUIPMENT", "LOCATION"}:
                raise AppError(422, "DOCUMENT_PARSE_FAILED", "RESOURCE_SHEET needs Name and Type (CAST/CREW/EQUIPMENT/LOCATION).")
            items.append({
                "name": name[:160], "type": type_, "subtype": _cell(row, "subtype")[:80],
                "email": _cell(row, "email") or None, "address": _cell(row, "address") or None,
                "default_cost": _integer(_cell(row, "cost", "default cost"), 0, 0, 1_000_000_000_000),
            })
        elif kind == "CALL_SHEET":
            scene = _cell(row, "scene", "scene number", "adegan")
            day = _cell(row, "day", "shooting day", "hari")
            start = _cell(row, "start", "start time", "call time", "call", "time", "waktu", "jam")
            if not scene or not day or not start:
                raise AppError(422, "DOCUMENT_PARSE_FAILED", "CALL_SHEET needs Day, Start and Scene columns.")
            items.append({"scene_number": scene[:20], "day": day[:20], "start": start[:40], "location": _cell(row, "location")[:160]})
        else:
            raise AppError(422, "DOCUMENT_UNSUPPORTED", "Unknown production document kind.")
    if kind == "CALL_SHEET":
        notes.append("Call Sheet times are shown for review only. Use Generate Initial Schedule until full imported schedule validation is implemented.")
    if kind == "SHOT_LIST":
        notes.append("Import SCRIPT scenes first. Unknown scene references will not be silently created.")
    notes.append("Existing records are preserved. Duplicate scene/shot/resource identifiers are skipped.")
    return {"document_id": document.id, "version_id": version.id, "type": kind,
            "count": len(items), "items": items, "notes": notes, "can_apply": kind != "CALL_SHEET"}


def apply(db: Session, project: Project, document_id: str, expected_version_id: str) -> dict[str, Any]:
    result = preview(db, project.id, document_id)
    if result["version_id"] != expected_version_id:
        raise AppError(409, "DOCUMENT_VERSION_OUTDATED", "Spreadsheet changed since preview. Review the new version first.")
    if not result["can_apply"]:
        raise AppError(409, "DOCUMENT_UNSUPPORTED", "This Call Sheet has not been verified for schedule import.")
    created = 0
    skipped = 0
    unknown: list[str] = []
    kind = result["type"]
    if kind == "SCRIPT":
        existing = {row.scene_number for row in db.scalars(select(Scene).where(Scene.project_id == project.id)).all()}
        for item in result["items"]:
            if item["scene_number"] in existing:
                skipped += 1
                continue
            db.add(Scene(project_id=project.id, **item))
            existing.add(item["scene_number"])
            created += 1
    elif kind == "SHOT_LIST":
        scene_map = {row.scene_number: row for row in db.scalars(select(Scene).where(Scene.project_id == project.id)).all()}
        ids = [row.id for row in scene_map.values()]
        existing = {(row.scene_id, row.shot_number) for row in db.scalars(select(Shot).where(Shot.scene_id.in_(ids))).all()} if ids else set()
        for item in result["items"]:
            number = item["scene_number"]
            scene = scene_map.get(number)
            if scene is None:
                unknown.append(number)
                continue
            key = (scene.id, item["shot_number"])
            if key in existing:
                skipped += 1
                continue
            values = {key: value for key, value in item.items() if key != "scene_number"}
            db.add(Shot(scene_id=scene.id, **values))
            existing.add(key)
            created += 1
    elif kind == "RESOURCE_SHEET":
        existing = {row.name.lower() for row in db.scalars(select(Resource).where(Resource.project_id == project.id)).all()}
        for item in result["items"]:
            if item["name"].lower() in existing:
                skipped += 1
                continue
            db.add(Resource(project_id=project.id, **item))
            existing.add(item["name"].lower())
            created += 1
    if created:
        project.production_state_version += 1
    db.commit()
    return {"created": created, "skipped": skipped, "unknown_scene_numbers": sorted(set(unknown)), "type": kind}
