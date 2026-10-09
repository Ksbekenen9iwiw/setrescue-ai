from __future__ import annotations

import hashlib
import io
import re
import uuid
import zipfile
from pathlib import Path

from fastapi import UploadFile
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.errors import AppError
from app.models import Document, DocumentVersion


ALLOWED_DOCUMENT_TYPES = {"SCRIPT", "SHOT_LIST", "CALL_SHEET", "RESOURCE_SHEET"}
ALLOWED_CONTENT_TYPES = {
    ".pdf": {"application/pdf", "application/octet-stream"},
    ".xlsx": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/octet-stream"},
}


async def read_limited(upload: UploadFile, limit: int) -> bytes:
    chunks: list[bytes] = []
    size = 0
    while chunk := await upload.read(1024 * 1024):
        size += len(chunk)
        if size > limit:
            raise AppError(413, "DOCUMENT_TOO_LARGE", "Uploaded document exceeds the allowed size.")
        chunks.append(chunk)
    return b"".join(chunks)


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(entry.file_size for entry in entries) > 50 * 1024 * 1024:
                raise AppError(422, "DOCUMENT_UNSUPPORTED", "XLSX archive expands beyond safe limits.")
            if "[Content_Types].xml" not in archive.namelist() or not any(name.startswith("xl/") for name in archive.namelist()):
                raise AppError(422, "DOCUMENT_UNSUPPORTED", "File is not a valid XLSX workbook.")
    except zipfile.BadZipFile as exc:
        raise AppError(422, "DOCUMENT_UNSUPPORTED", "File is not a valid XLSX workbook.") from exc


def validate_and_extract(filename: str, content_type: str | None, content: bytes) -> tuple[str, str]:
    extension = Path(filename or "").suffix.lower()
    if extension not in ALLOWED_CONTENT_TYPES:
        raise AppError(415, "DOCUMENT_UNSUPPORTED", "Only text-based PDF and XLSX documents are supported.")
    if content_type and content_type.lower().split(";", 1)[0] not in ALLOWED_CONTENT_TYPES[extension]:
        raise AppError(415, "DOCUMENT_UNSUPPORTED", "Document MIME type does not match an allowed format.")
    if extension == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise AppError(422, "DOCUMENT_UNSUPPORTED", "File signature does not match PDF.")
        try:
            reader = PdfReader(io.BytesIO(content), strict=True)
            if len(reader.pages) > 500:
                raise AppError(422, "DOCUMENT_UNSUPPORTED", "PDF contains too many pages.")
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        except AppError:
            raise
        except Exception as exc:
            raise AppError(422, "DOCUMENT_PARSE_FAILED", "PDF text could not be extracted.") from exc
        if not text.strip():
            raise AppError(422, "DOCUMENT_PARSE_FAILED", "PDF contains no extractable text; scanned PDFs are not supported.")
        return extension, text[:2_000_000]

    if not content.startswith(b"PK\x03\x04"):
        raise AppError(422, "DOCUMENT_UNSUPPORTED", "File signature does not match XLSX.")
    _validate_xlsx_archive(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
        rows: list[str] = []
        for sheet in workbook.worksheets[:50]:
            rows.append(f"[Sheet: {sheet.title}]")
            for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                if row_index >= 10_000:
                    break
                cells = [str(value)[:2000] for value in row if value is not None]
                if cells:
                    rows.append(" | ".join(cells))
        workbook.close()
        text = "\n".join(rows)
    except Exception as exc:
        raise AppError(422, "DOCUMENT_PARSE_FAILED", "XLSX content could not be extracted.") from exc
    return extension, text[:2_000_000]


async def save_document(db: Session, project_id: str, document_type: str, upload: UploadFile) -> DocumentVersion:
    if document_type not in ALLOWED_DOCUMENT_TYPES:
        raise AppError(422, "DOCUMENT_UNSUPPORTED", "Unknown document type.")
    extension = Path(upload.filename or "").suffix.lower()
    limit = get_settings().max_pdf_bytes if extension == ".pdf" else get_settings().max_xlsx_bytes
    content = await read_limited(upload, limit)
    file_type, text = validate_and_extract(upload.filename or "document", upload.content_type, content)
    digest = hashlib.sha256(content).hexdigest()

    document = db.scalar(select(Document).where(Document.project_id == project_id, Document.type == document_type))
    if document:
        document.current_version += 1
    else:
        document = Document(project_id=project_id, type=document_type, current_version=1)
        db.add(document)
        db.flush()

    project_dir = get_settings().data_dir / "projects" / project_id / "documents"
    project_dir.mkdir(parents=True, exist_ok=True)
    generated_name = f"{uuid.uuid4().hex}{file_type}"
    destination = (project_dir / generated_name).resolve()
    if project_dir.resolve() not in destination.parents:
        raise AppError(500, "DOCUMENT_STORE_FAILED", "Document storage path is invalid.")
    destination.write_bytes(content)
    clean_original = re.sub(r"[^A-Za-z0-9._ -]", "_", Path(upload.filename or "document").name)[:255]
    version = DocumentVersion(
        document_id=document.id, version_number=document.current_version,
        original_filename=clean_original, local_path=str(destination), file_type=file_type,
        file_hash=digest, extracted_text=text,
    )
    db.add(version)
    # AI inputs are content-hashed, so a modified document produces a different cache key.
    # Do not evict all cached analyses when an unchanged document is reuploaded.
    db.commit()
    db.refresh(version)
    return version
